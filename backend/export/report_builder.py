"""Build formatted Excel and PowerPoint reports from DDA attribution results."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt


_HEADER_FILL = PatternFill(start_color="1e293b", end_color="1e293b", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="ffffff", size=10)
_THIN_BORDER = Border(
    bottom=Side(style="thin", color="334155"),
)
_PCT_FMT = "0.0%"
_NUM_FMT = "#,##0"
_DEC_FMT = "0.00"


def _auto_widths(ws) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            val = str(cell.value) if cell.value is not None else ""
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


def _write_header(ws, row: int, values: list[str]) -> None:
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = _THIN_BORDER


def _write_title(ws, campaign_name: str, run_date: str) -> None:
    ws.cell(row=1, column=1, value=f"Time's Hub — {campaign_name}").font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value=f"Rapor Tarihi: {run_date}").font = Font(size=9, color="64748b")


def _build_summary_sheet(
    wb: Workbook, result: dict, campaign_name: str, run_date: str, objective: str = "lead"
) -> None:
    ws = wb.active
    ws.title = "Özet"
    _write_title(ws, campaign_name, run_date)

    stats = result.get("journey_stats", {})
    lead_label = "Toplam Lead" if objective == "lead" else "Dönüşüm Yapan"
    _write_header(ws, 4, ["Metrik", "Değer"])

    rows = [
        ("Toplam Yolculuk", stats.get("total_journeys", 0), _NUM_FMT),
        (lead_label, stats.get("converted", 0), _NUM_FMT),
        ("Dönüşüm Oranı", stats.get("conversion_rate", 0), _PCT_FMT),
        ("Ort. Temas Noktası", stats.get("avg_path_length", stats.get("avg_touchpoints", 0)), _DEC_FMT),
        ("Tek Temas %", stats.get("single_touch_pct", 0), _PCT_FMT),
        ("Çoklu Temas %", stats.get("multi_touch_pct", 0), _PCT_FMT),
    ]

    # Revenue mode: append revenue totals when available
    total_revenue = result.get("bq_summary", {}).get("total_revenue", 0)
    if objective == "revenue" and total_revenue:
        converted = stats.get("converted", 0)
        aov = total_revenue / converted if converted else 0
        rows.append(("Toplam Gelir", total_revenue, _NUM_FMT))
        rows.append(("AOV (Ort. Sipariş Değeri)", aov, _NUM_FMT))

    for i, (label, val, fmt) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=val)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")

    markov = result.get("markov", {})
    conv_prob = markov.get("conversion_probability", 0)
    if conv_prob:
        r = len(rows) + 6
        ws.cell(row=r, column=1, value="Markov Dönüşüm Olasılığı")
        c = ws.cell(row=r, column=2, value=conv_prob)
        c.number_format = _PCT_FMT

    _auto_widths(ws)


def _build_attribution_sheet(wb: Workbook, result: dict, objective: str = "lead") -> None:
    ws = wb.create_sheet("Kanal Atfetme")
    hybrid = result.get("hybrid_attribution", {})
    markov_w = result.get("markov", {}).get("attribution_weights", {})
    shapley = result.get("shapley_dda", {})
    removal = result.get("markov", {}).get("removal_effects", {})

    # Attributed total: leads (converted) for lead mode, revenue for revenue mode
    stats = result.get("journey_stats", {})
    total_leads = stats.get("converted", 0)
    total_revenue = result.get("bq_summary", {}).get("total_revenue", 0)
    if objective == "revenue" and total_revenue:
        attr_label, attr_total, attr_fmt = "Atf. Gelir (₺)", total_revenue, _NUM_FMT
    else:
        attr_label, attr_total, attr_fmt = "Atf. Lead", total_leads, _NUM_FMT

    _write_header(
        ws, 1,
        ["Kanal", "DDA Katkı (%)", "Markov (%)", "Shapley (%)", "Kaldırma Etkisi", attr_label],
    )

    channels = sorted(hybrid.keys(), key=lambda ch: -hybrid.get(ch, 0))
    for i, ch in enumerate(channels, 2):
        ws.cell(row=i, column=1, value=ch)
        for col, src in [(2, hybrid), (3, markov_w), (4, shapley)]:
            c = ws.cell(row=i, column=col, value=src.get(ch, 0))
            c.number_format = _PCT_FMT
            c.alignment = Alignment(horizontal="right")
        c = ws.cell(row=i, column=5, value=removal.get(ch, 0))
        c.number_format = _DEC_FMT
        c.alignment = Alignment(horizontal="right")
        # Attributed value = total × hybrid weight
        c = ws.cell(row=i, column=6, value=round(attr_total * hybrid.get(ch, 0), 1))
        c.number_format = attr_fmt
        c.alignment = Alignment(horizontal="right")

    _auto_widths(ws)


def _build_assist_sheet(wb: Workbook, result: dict) -> None:
    ws = wb.create_sheet("Asist Raporu")
    report = result.get("assist_report", [])
    _write_header(ws, 1, ["Kanal", "Son Temas", "İlk Temas", "Asist", "Asist Oranı", "Toplam Katılım"])

    for i, row in enumerate(report, 2):
        ws.cell(row=i, column=1, value=row.get("channel", ""))
        ws.cell(row=i, column=2, value=row.get("last_touch", 0)).number_format = _NUM_FMT
        ws.cell(row=i, column=3, value=row.get("first_touch", 0)).number_format = _NUM_FMT
        ws.cell(row=i, column=4, value=row.get("assists", 0)).number_format = _NUM_FMT
        ws.cell(row=i, column=5, value=row.get("assist_ratio", 0)).number_format = _PCT_FMT
        ws.cell(row=i, column=6, value=row.get("total_involvement", 0)).number_format = _NUM_FMT

    _auto_widths(ws)


def _build_paths_sheet(wb: Workbook, result: dict) -> None:
    paths = result.get("top_paths", [])
    if not paths:
        return
    ws = wb.create_sheet("Dönüşüm Yolları")
    _write_header(ws, 1, ["#", "Yol", "Sayı", "Dönüşüm Oranı"])

    for i, p in enumerate(paths[:20], 2):
        ws.cell(row=i, column=1, value=i - 1)
        path_parts = p.get("path", [])
        if isinstance(path_parts, list):
            path_str = " → ".join(str(s) for s in path_parts)
        else:
            path_str = str(path_parts)
        ws.cell(row=i, column=2, value=path_str)
        ws.cell(row=i, column=3, value=p.get("count", p.get("total", 0))).number_format = _NUM_FMT
        ws.cell(row=i, column=4, value=p.get("conversion_rate", p.get("rate", 0))).number_format = _PCT_FMT

    _auto_widths(ws)


def _build_insights_sheet(wb: Workbook, result: dict) -> None:
    insights = result.get("insights", [])
    if not insights:
        return
    ws = wb.create_sheet("Çıkarımlar")

    type_fills = {
        "warning": PatternFill(start_color="451a03", end_color="451a03", fill_type="solid"),
        "success": PatternFill(start_color="052e16", end_color="052e16", fill_type="solid"),
        "info": PatternFill(start_color="0c1a3d", end_color="0c1a3d", fill_type="solid"),
    }
    type_fonts = {
        "warning": Font(color="fbbf24", size=10),
        "success": Font(color="34d399", size=10),
        "info": Font(color="60a5fa", size=10),
    }

    _write_header(ws, 1, ["Tip", "Kategori", "Çıkarım"])

    for i, ins in enumerate(insights, 2):
        t = ins.get("type", "info")
        ws.cell(row=i, column=1, value=t.capitalize())
        ws.cell(row=i, column=2, value=ins.get("category", ""))
        c = ws.cell(row=i, column=3, value=ins.get("text", ""))
        fill = type_fills.get(t)
        font = type_fonts.get(t)
        if fill:
            for col in range(1, 4):
                ws.cell(row=i, column=col).fill = fill
        if font:
            c.font = font

    _auto_widths(ws)


def build_dda_report(
    result: dict, campaign_name: str = "", run_date: str = "", objective: str = "lead"
) -> Workbook:
    """Build a multi-sheet Excel workbook from a DDA result snapshot.

    objective ("lead" | "revenue") drives whether the attributed-value column
    shows leads or revenue, and whether revenue totals appear in the summary.
    """
    wb = Workbook()
    _build_summary_sheet(wb, result, campaign_name, run_date, objective)
    _build_attribution_sheet(wb, result, objective)
    _build_assist_sheet(wb, result)
    _build_paths_sheet(wb, result)
    _build_insights_sheet(wb, result)
    return wb


def workbook_to_bytes(wb: Workbook) -> BytesIO:
    """Serialize a Workbook to a BytesIO buffer ready for streaming."""
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# --------------- PowerPoint Report ---------------


_BRAND_DARK = RGBColor(0x1E, 0x29, 0x3B)
_BRAND_ACCENT = RGBColor(0xF9, 0x73, 0x16)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_GRAY = RGBColor(0x94, 0xA3, 0xB8)

_CHANNEL_COLORS = {
    "meta": RGBColor(0x38, 0x82, 0xF6),
    "google": RGBColor(0x22, 0xC5, 0x5E),
    "tiktok": RGBColor(0xA8, 0x55, 0xF7),
    "linkedin": RGBColor(0x06, 0xB6, 0xD4),
    "dv360": RGBColor(0xF9, 0x73, 0x16),
    "youtube": RGBColor(0xEF, 0x44, 0x44),
}


def _slide_bg(slide, color=_BRAND_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text_box(slide, left, top, width, height, text, font_size=12,
                  color=_WHITE, bold=False, alignment=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.alignment = alignment
    return txBox


def _slide_title(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_bg(slide)
    _add_text_box(slide, Inches(0.8), Inches(2.0), Inches(8.4), Inches(1.0),
                  title, font_size=28, bold=True, alignment=PP_ALIGN.CENTER)
    if subtitle:
        _add_text_box(slide, Inches(0.8), Inches(3.2), Inches(8.4), Inches(0.8),
                      subtitle, font_size=14, color=_GRAY, alignment=PP_ALIGN.CENTER)
    return slide


def _slide_summary(prs, result, objective="lead"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_bg(slide)
    _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.6),
                  "Yolculuk Özeti", font_size=22, bold=True, color=_BRAND_ACCENT)

    stats = result.get("journey_stats", {})
    lead_label = "Toplam Lead" if objective == "lead" else "Dönüşüm"
    kpis = [
        ("Toplam Yolculuk", f"{stats.get('total_journeys', 0):,}"),
        (lead_label, f"{stats.get('converted', 0):,}"),
        ("Dönüşüm Oranı", f"%{stats.get('conversion_rate', 0) * 100:.1f}"),
        ("Ort. Temas Noktası", f"{stats.get('avg_path_length', 0):.1f}"),
    ]

    for i, (label, value) in enumerate(kpis):
        left = Inches(0.5 + i * 2.4)
        _add_text_box(slide, left, Inches(1.4), Inches(2.2), Inches(0.5),
                      value, font_size=32, bold=True, color=_WHITE,
                      alignment=PP_ALIGN.CENTER)
        _add_text_box(slide, left, Inches(2.1), Inches(2.2), Inches(0.4),
                      label, font_size=11, color=_GRAY,
                      alignment=PP_ALIGN.CENTER)

    markov = result.get("markov", {})
    conv_prob = markov.get("conversion_probability", 0)
    if conv_prob:
        _add_text_box(slide, Inches(0.5), Inches(3.5), Inches(9), Inches(0.4),
                      f"Markov Dönüşüm Olasılığı: %{conv_prob * 100:.1f}",
                      font_size=12, color=_GRAY)
    return slide


def _slide_attribution(prs, result, objective="lead"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_bg(slide)
    _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.6),
                  "DDA Kanal Attribution", font_size=22, bold=True, color=_BRAND_ACCENT)

    hybrid = result.get("hybrid_attribution", {})
    if not hybrid:
        return slide

    channels = sorted(hybrid.keys(), key=lambda ch: -hybrid.get(ch, 0))

    stats = result.get("journey_stats", {})
    total_leads = stats.get("converted", 0)
    total_revenue = result.get("bq_summary", {}).get("total_revenue", 0)

    bar_max_width = Inches(5.5)
    bar_height = Emu(int(Inches(0.35)))
    y_start = Inches(1.3)
    row_gap = Inches(0.55)

    for i, ch in enumerate(channels[:8]):
        y = y_start + i * row_gap
        weight = hybrid.get(ch, 0)
        bar_width = int(bar_max_width * weight) if weight > 0 else Emu(1)

        _add_text_box(slide, Inches(0.5), y, Inches(1.3), bar_height,
                      ch.capitalize(), font_size=11, bold=True, color=_WHITE)

        bar = slide.shapes.add_shape(
            1, Inches(2.0), y, bar_width, bar_height,
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = _CHANNEL_COLORS.get(ch, _BRAND_ACCENT)
        bar.line.fill.background()

        pct_text = f"%{weight * 100:.1f}"
        if objective == "revenue" and total_revenue:
            attr_val = f" — ₺{total_revenue * weight:,.0f}"
        else:
            attr_val = f" — {total_leads * weight:.0f} lead" if total_leads else ""

        _add_text_box(slide, Inches(2.0) + bar_width + Emu(int(Inches(0.15))),
                      y, Inches(3), bar_height,
                      pct_text + attr_val, font_size=10, color=_GRAY)

    _add_text_box(slide, Inches(0.5), Inches(6.8), Inches(9), Inches(0.4),
                  "DDA: Markov Chain (%65) + Shapley Value (%35) ensemble — gerçek yolculuk verisinden hesaplanmıştır.",
                  font_size=9, color=_GRAY, alignment=PP_ALIGN.CENTER)
    return slide


def _slide_assist(prs, result):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_bg(slide)
    _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.6),
                  "Asist Dönüşüm Raporu", font_size=22, bold=True, color=_BRAND_ACCENT)

    report = result.get("assist_report", [])
    if not report:
        _add_text_box(slide, Inches(0.5), Inches(2), Inches(9), Inches(0.5),
                      "Asist verisi bulunamadı.", font_size=14, color=_GRAY)
        return slide

    headers = ["Kanal", "Son Temas", "İlk Temas", "Asist", "Asist Oranı", "Rol"]
    col_widths = [Inches(1.5), Inches(1.2), Inches(1.2), Inches(1.0), Inches(1.3), Inches(1.5)]
    x_start = Inches(0.8)
    y_header = Inches(1.3)

    x = x_start
    for h, w in zip(headers, col_widths):
        _add_text_box(slide, x, y_header, w, Inches(0.35),
                      h, font_size=10, bold=True, color=_BRAND_ACCENT,
                      alignment=PP_ALIGN.CENTER)
        x += w

    for i, r in enumerate(report[:8]):
        y = y_header + Inches(0.4) + i * Inches(0.4)
        vals = [
            r.get("channel", ""),
            str(r.get("last_touch", 0)),
            str(r.get("first_touch", 0)),
            str(r.get("assists", 0)),
            f"%{r.get('assist_ratio', 0) * 100:.0f}",
            r.get("channel_role", "Hibrit"),
        ]
        x = x_start
        for val, w in zip(vals, col_widths):
            _add_text_box(slide, x, y, w, Inches(0.35),
                          val, font_size=10, color=_WHITE,
                          alignment=PP_ALIGN.CENTER)
            x += w

    return slide


def _slide_insights(prs, result):
    insights = result.get("insights", [])
    if not insights:
        return None

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _slide_bg(slide)
    _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.6),
                  "Otomatik Çıkarımlar", font_size=22, bold=True, color=_BRAND_ACCENT)

    type_colors = {
        "warning": RGBColor(0xFB, 0xBF, 0x24),
        "success": RGBColor(0x34, 0xD3, 0x99),
        "info": RGBColor(0x60, 0xA5, 0xFA),
    }

    y = Inches(1.3)
    for ins in insights[:10]:
        t = ins.get("type", "info")
        icon = {"warning": "!", "success": "+", "info": "i"}.get(t, "·")
        color = type_colors.get(t, _WHITE)
        text = f"[{icon}] {ins.get('text', '')}"
        _add_text_box(slide, Inches(0.8), y, Inches(8.4), Inches(0.5),
                      text, font_size=11, color=color)
        y += Inches(0.5)

    return slide


def build_dda_pptx(
    result: dict, campaign_name: str = "", run_date: str = "", objective: str = "lead"
) -> Presentation:
    """Build a PowerPoint presentation from DDA attribution results."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    _slide_title(prs, f"Attribution Raporu — {campaign_name}",
                 f"DDA Analiz Sonuçları  •  {run_date}")
    _slide_summary(prs, result, objective)
    _slide_attribution(prs, result, objective)
    _slide_assist(prs, result)
    _slide_insights(prs, result)

    return prs


def pptx_to_bytes(prs: Presentation) -> BytesIO:
    """Serialize a Presentation to BytesIO buffer."""
    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
