"""The "traffic" (awareness / site-visit) campaign objective."""

import io

from openpyxl import load_workbook

from backend.models.simulation import simulate_budget


def _client(client, headers, objective="traffic"):
    res = client.post("/api/clients", json={"name": f"Trafik Co {objective}", "year": 2026, "objective": objective},
                      headers=headers)
    assert res.status_code == 200, res.text
    return res.json()


def test_traffic_client_and_inherited_campaign(client, auth_headers):
    c = _client(client, auth_headers)
    assert c["objective"] == "traffic"
    camp = client.post(f"/api/clients/{c['id']}/campaigns", json={"name": "Lansman", "budget": 15_000_000},
                       headers=auth_headers)
    assert camp.status_code == 200, camp.text
    assert camp.json()["objective"] == "traffic"  # inherited from the client


def test_campaign_can_switch_to_traffic(client, auth_headers):
    c = _client(client, auth_headers, objective="lead")
    camp = client.post(f"/api/clients/{c['id']}/campaigns", json={"name": "X"}, headers=auth_headers).json()
    res = client.patch(f"/api/campaigns/{camp['id']}", json={"objective": "traffic"}, headers=auth_headers)
    assert res.status_code == 200 and res.json()["objective"] == "traffic"


def test_unknown_objective_rejected_with_turkish_message(client, auth_headers):
    res = client.post("/api/clients", json={"name": "Bad", "year": 2026, "objective": "reach"}, headers=auth_headers)
    assert res.status_code == 400
    assert "traffic" in res.json()["detail"] and "biri olmalı" in res.json()["detail"]


def test_traffic_simulation_is_count_based_without_revenue():
    out = simulate_budget(
        channel_spends={"google/cpc": 100_000, "meta/paid": 50_000},
        dda_weights={"google/cpc": 0.6, "meta/paid": 0.4},
        total_revenue=0.0,
        total_conversions=1_000,  # qualified visits
        objective="traffic",
        lead_value=500,  # ignored outside lead mode
    )
    assert out["objective"] == "traffic" and out["primary_metric"] == "visits"
    google = out["current"]["channels"]["google/cpc"]
    assert google["attributed_conversions"] == 600.0
    assert google["cpa"] == round(100_000 / 600, 2)  # cost per visit
    assert "value_roas" not in google and "blended_value_roas" not in out["current"]


def test_budget_endpoint_accepts_traffic_without_revenue(client, auth_headers):
    res = client.post("/api/simulation/budget", json={
        "channel_spends": {"google/cpc": 100_000}, "dda_weights": {"google/cpc": 1.0},
        "total_conversions": 500, "objective": "traffic",
    }, headers=auth_headers)
    assert res.status_code == 200, res.text
    assert res.json()["primary_metric"] == "visits"


def test_revenue_mode_still_requires_revenue(client, auth_headers):
    res = client.post("/api/simulation/budget", json={
        "channel_spends": {"google/cpc": 100_000}, "dda_weights": {"google/cpc": 1.0},
        "total_conversions": 500, "objective": "revenue",
    }, headers=auth_headers)
    assert res.status_code == 400


def test_exports_use_visit_wording(client, auth_headers, sample_journeys_csv):
    c = _client(client, auth_headers)
    camp = client.post(f"/api/clients/{c['id']}/campaigns", json={"name": "Rapor"}, headers=auth_headers).json()
    run = client.post("/api/dda/run-from-csv", params={"campaign_id": camp["id"]},
                      files={"file": ("j.csv", sample_journeys_csv, "text/csv")}, headers=auth_headers)
    assert run.status_code == 200, run.text

    xlsx = client.get("/api/export/dda-report", params={"campaign_id": camp["id"]}, headers=auth_headers)
    assert xlsx.status_code == 200 and "trafik" in xlsx.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(xlsx.content))
    cells = {str(cell.value) for ws in wb for row in ws.iter_rows() for cell in row if cell.value}
    assert "Toplam Ziyaret" in cells and "Atf. Ziyaret" in cells
    assert "Toplam Lead" not in cells

    pptx = client.get("/api/export/dda-pptx", params={"campaign_id": camp["id"]}, headers=auth_headers)
    assert pptx.status_code == 200 and "trafik" in pptx.headers["content-disposition"]
