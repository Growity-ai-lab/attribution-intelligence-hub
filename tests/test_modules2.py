"""Unit tests for simulation, report_builder, and mmm_fit modules."""

import os

os.environ.setdefault("AUTH_SECRET_KEY", "test-secret-key-for-ci")
os.environ.setdefault("AUTH_ADMIN_PASSWORD", "test-password-for-ci")

import pytest
from io import BytesIO
from types import SimpleNamespace

from backend.models.simulation import simulate_budget, generate_budget_recommendations
from backend.export.report_builder import (
    build_dda_pptx,
    build_dda_report,
    pptx_to_bytes,
    workbook_to_bytes,
)
from backend.models.mmm_fit import fit_mmm, compute_data_hash


# ── simulate_budget ──────────────────────────────────────────────────


class TestSimulateBudget:
    def test_output_structure(self):
        result = simulate_budget(
            channel_spends={"meta": 100_000, "google": 50_000},
            dda_weights={"meta": 0.6, "google": 0.4},
            total_revenue=500_000,
            total_conversions=100,
        )
        assert "current" in result
        assert "recommendations" in result
        cur = result["current"]
        assert cur["total_spend"] == 150_000
        assert cur["total_revenue"] == 500_000
        assert "meta" in cur["channels"]
        assert "google" in cur["channels"]

    def test_attributed_revenue(self):
        result = simulate_budget(
            channel_spends={"meta": 100_000},
            dda_weights={"meta": 1.0},
            total_revenue=500_000,
            total_conversions=100,
        )
        ch = result["current"]["channels"]["meta"]
        assert ch["attributed_revenue"] == 500_000
        assert ch["attributed_conversions"] == 100

    def test_zero_spend_roas_none(self):
        result = simulate_budget(
            channel_spends={"meta": 0},
            dda_weights={"meta": 1.0},
            total_revenue=500_000,
            total_conversions=100,
        )
        ch = result["current"]["channels"]["meta"]
        assert ch["roas"] is None

    def test_zero_conversions_cpa_none(self):
        result = simulate_budget(
            channel_spends={"meta": 100_000},
            dda_weights={"meta": 1.0},
            total_revenue=500_000,
            total_conversions=0,
        )
        ch = result["current"]["channels"]["meta"]
        assert ch["cpa"] is None

    def test_scenario_projection(self):
        result = simulate_budget(
            channel_spends={"meta": 100_000, "google": 50_000},
            dda_weights={"meta": 0.6, "google": 0.4},
            total_revenue=500_000,
            total_conversions=100,
            scenario_spends={"meta": 200_000, "google": 50_000},
        )
        assert "scenario" in result

    def test_empty_weights(self):
        result = simulate_budget(
            channel_spends={"meta": 100_000},
            dda_weights={},
            total_revenue=500_000,
            total_conversions=100,
        )
        assert result["current"]["channels"] == {} or all(
            ch["weight"] == 0 for ch in result["current"]["channels"].values()
        )

    def test_blended_roas(self):
        result = simulate_budget(
            channel_spends={"meta": 100_000, "google": 50_000},
            dda_weights={"meta": 0.6, "google": 0.4},
            total_revenue=300_000,
            total_conversions=100,
        )
        roas = result["current"]["blended_roas"]
        assert roas is not None
        assert abs(roas - 2.0) < 0.01


# ── generate_budget_recommendations ──────────────────────────────────


class TestBudgetRecommendations:
    def _channels(self, roas_values):
        return {
            f"ch{i}": {
                "spend": 100_000,
                "attributed_revenue": 100_000 * r,
                "attributed_conversions": 10,
                "weight": 0.2,
                "roas": r,
                "cpa": 10_000,
                "organic": False,
            }
            for i, r in enumerate(roas_values)
        }

    def test_high_roas_gets_increase(self):
        channels = self._channels([5.0, 1.0, 1.0])
        recs = generate_budget_recommendations(channels, total_spend=300_000)
        high_rec = next(r for r in recs if r["channel"] == "ch0")
        assert high_rec["action"] == "artır" or high_rec["action"] == "artir"

    def test_low_roas_gets_decrease(self):
        channels = self._channels([0.2, 3.0, 3.0])
        recs = generate_budget_recommendations(channels, total_spend=300_000)
        low_rec = next(r for r in recs if r["channel"] == "ch0")
        assert low_rec["action"] == "azalt"

    def test_empty_channels(self):
        recs = generate_budget_recommendations({})
        assert recs == []

    def test_organic_channel_gets_special_action(self):
        channels = {
            "organic_search": {
                "spend": 0,
                "attributed_revenue": 50_000,
                "attributed_conversions": 20,
                "weight": 0.15,
                "roas": None,
                "cpa": None,
                "organic": True,
            }
        }
        recs = generate_budget_recommendations(channels)
        if recs:
            assert recs[0]["action"] in ("değerlendirmeli", "degerlendirmeli", "koru")


# ── build_dda_report ─────────────────────────────────────────────────


def _dda_result():
    return {
        "journey_stats": {
            "total_journeys": 100,
            "converted": 50,
            "conversion_rate": 0.5,
            "avg_path_length": 2.5,
            "single_touch_pct": 0.3,
            "multi_touch_pct": 0.7,
        },
        "hybrid_attribution": {"meta": 0.4, "google": 0.35, "tiktok": 0.25},
        "markov": {
            "attribution_weights": {"meta": 0.45, "google": 0.3, "tiktok": 0.25},
            "removal_effects": {"meta": 0.5, "google": 0.3, "tiktok": 0.2},
            "conversion_probability": 0.42,
        },
        "shapley_dda": {"meta": 0.35, "google": 0.4, "tiktok": 0.25},
        "assist_report": [
            {
                "channel": "meta",
                "first_touch": 30,
                "last_touch": 15,
                "assists": 20,
                "assist_ratio": 0.57,
                "total_involvement": 65,
            },
        ],
        "top_paths": [
            {"path": ["meta", "google", "form"], "count": 25, "conversion_rate": 0.6},
            {"path": ["tiktok", "meta", "form"], "count": 15, "rate": 0.4},
        ],
        "insights": [
            {"type": "warning", "icon": "⚠", "text": "Meta yoğunlaşması yüksek", "category": "concentration"},
            {"type": "success", "icon": "✓", "text": "Dönüşüm oranı iyi", "category": "conversion"},
        ],
    }


class TestBuildDDAReport:
    def test_returns_workbook(self):
        wb = build_dda_report(_dda_result(), "Test Campaign", "2026-06-01")
        from openpyxl import Workbook

        assert isinstance(wb, Workbook)

    def test_sheet_names(self):
        wb = build_dda_report(_dda_result(), "Test Campaign", "2026-06-01")
        names = wb.sheetnames
        assert "Özet" in names
        assert "Kanal Atfetme" in names
        assert "Asist Raporu" in names

    def test_paths_sheet_present(self):
        wb = build_dda_report(_dda_result(), "Test Campaign", "2026-06-01")
        assert "Dönüşüm Yolları" in wb.sheetnames

    def test_insights_sheet_present(self):
        wb = build_dda_report(_dda_result(), "Test Campaign", "2026-06-01")
        assert "Çıkarımlar" in wb.sheetnames

    def test_no_paths_sheet_when_empty(self):
        result = _dda_result()
        result["top_paths"] = []
        wb = build_dda_report(result, "Test", "2026-06-01")
        assert "Dönüşüm Yolları" not in wb.sheetnames

    def test_no_insights_sheet_when_empty(self):
        result = _dda_result()
        result["insights"] = []
        wb = build_dda_report(result, "Test", "2026-06-01")
        assert "Çıkarımlar" not in wb.sheetnames

    def test_attribution_channels_match(self):
        result = _dda_result()
        wb = build_dda_report(result, "Test", "2026-06-01")
        ws = wb["Kanal Atfetme"]
        channel_cells = []
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                channel_cells.append(row[0])
        for ch in result["hybrid_attribution"]:
            assert ch in channel_cells

    def test_path_joined_with_arrow(self):
        result = _dda_result()
        wb = build_dda_report(result, "Test", "2026-06-01")
        ws = wb["Dönüşüm Yolları"]
        found_arrow = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] and "→" in str(row[1]):
                found_arrow = True
                break
        assert found_arrow

    def test_empty_result_no_crash(self):
        wb = build_dda_report({}, "Empty", "2026-06-01")
        assert len(wb.sheetnames) >= 1


class TestWorkbookToBytes:
    def test_returns_bytesio(self):
        wb = build_dda_report(_dda_result(), "Test", "2026-06-01")
        buf = workbook_to_bytes(wb)
        assert isinstance(buf, BytesIO)

    def test_non_empty(self):
        wb = build_dda_report(_dda_result(), "Test", "2026-06-01")
        buf = workbook_to_bytes(wb)
        assert len(buf.getvalue()) > 0

    def test_buffer_at_start(self):
        wb = build_dda_report(_dda_result(), "Test", "2026-06-01")
        buf = workbook_to_bytes(wb)
        assert buf.tell() == 0


# ── fit_mmm ──────────────────────────────────────────────────────────


class TestFitMMM:
    @pytest.fixture()
    def sample_data(self):
        import numpy as np

        np.random.seed(42)
        weeks = 10
        spends = {
            "meta": [float(x) for x in np.random.uniform(50_000, 200_000, weeks)],
            "google": [float(x) for x in np.random.uniform(20_000, 80_000, weeks)],
        }
        leads = [float(x) for x in np.random.uniform(50, 200, weeks)]
        return spends, leads

    def test_output_structure(self, sample_data):
        spends, leads = sample_data
        result = fit_mmm(spends, leads)
        assert "channels" in result
        assert "params" in result
        assert "baseline" in result
        assert "fit_quality" in result
        assert "residuals" in result
        assert "predicted" in result
        assert "actual" in result

    def test_channels_list(self, sample_data):
        spends, leads = sample_data
        result = fit_mmm(spends, leads)
        assert set(result["channels"]) == {"meta", "google"}

    def test_param_keys(self, sample_data):
        spends, leads = sample_data
        result = fit_mmm(spends, leads)
        for ch in result["channels"]:
            p = result["params"][ch]
            assert "decay" in p
            assert "alpha" in p
            assert "gamma" in p
            assert "max_lift" in p

    def test_fit_quality_keys(self, sample_data):
        spends, leads = sample_data
        result = fit_mmm(spends, leads)
        fq = result["fit_quality"]
        assert "rmse" in fq
        assert "mape" in fq
        assert "r2" in fq
        assert "n_obs" in fq
        assert "converged" in fq
        assert fq["n_obs"] == 10

    def test_residuals_length(self, sample_data):
        spends, leads = sample_data
        result = fit_mmm(spends, leads)
        assert len(result["residuals"]) == 10
        assert len(result["predicted"]) == 10
        assert len(result["actual"]) == 10

    def test_non_negative_baseline(self, sample_data):
        spends, leads = sample_data
        result = fit_mmm(spends, leads)
        assert result["baseline"] >= 0

    def test_decay_within_bounds(self, sample_data):
        spends, leads = sample_data
        result = fit_mmm(spends, leads)
        for ch in result["channels"]:
            d = result["params"][ch]["decay"]
            assert 0.0 <= d <= 0.99

    def test_raises_on_empty_channels(self):
        with pytest.raises(ValueError, match="[Nn]o channels"):
            fit_mmm({}, [100, 200, 300, 400])

    def test_raises_on_short_data(self):
        with pytest.raises(ValueError, match="4"):
            fit_mmm({"meta": [100, 200, 300]}, [50, 60, 70])

    def test_raises_on_mismatched_lengths(self):
        with pytest.raises(ValueError):
            fit_mmm({"meta": [100, 200, 300, 400]}, [50, 60])

    def test_single_channel(self):
        import numpy as np

        np.random.seed(99)
        spends = {"meta": [float(x) for x in np.random.uniform(50_000, 200_000, 8)]}
        leads = [float(x) for x in np.random.uniform(50, 200, 8)]
        result = fit_mmm(spends, leads)
        assert result["channels"] == ["meta"]


# ── compute_data_hash ────────────────────────────────────────────────


class TestComputeDataHash:
    def _rows(self, data):
        return [
            SimpleNamespace(week=w, channel=c, spend=s, leads=l) for w, c, s, l in data
        ]

    def test_consistency(self):
        rows = self._rows([("W01", "meta", 100.0, 10)])
        h1 = compute_data_hash(rows)
        h2 = compute_data_hash(rows)
        assert h1 == h2

    def test_different_data_different_hash(self):
        r1 = self._rows([("W01", "meta", 100.0, 10)])
        r2 = self._rows([("W01", "meta", 200.0, 10)])
        assert compute_data_hash(r1) != compute_data_hash(r2)

    def test_empty_list(self):
        h = compute_data_hash([])
        assert isinstance(h, str)
        assert len(h) == 32

    def test_order_independent(self):
        r1 = self._rows([("W01", "meta", 100.0, 10), ("W02", "google", 200.0, 20)])
        r2 = self._rows([("W02", "google", 200.0, 20), ("W01", "meta", 100.0, 10)])
        assert compute_data_hash(r1) == compute_data_hash(r2)

    def test_md5_hex_format(self):
        rows = self._rows([("W01", "meta", 100.0, 10)])
        h = compute_data_hash(rows)
        assert len(h) == 32
        assert all(c in "0123456789abcdef" for c in h)


# ── PPTX export ──────────────────────────────────────────────────────

def _sample_dda_result():
    return {
        "journey_stats": {
            "total_journeys": 100,
            "converted": 40,
            "conversion_rate": 0.4,
            "avg_path_length": 2.3,
        },
        "hybrid_attribution": {"meta": 0.45, "google": 0.35, "tiktok": 0.20},
        "markov": {
            "conversion_probability": 0.38,
            "removal_effects": {"meta": 0.52, "google": 0.30, "tiktok": 0.18},
            "attribution_weights": {"meta": 0.52, "google": 0.30, "tiktok": 0.18},
            "warnings": [],
        },
        "shapley_dda": {"meta": 0.40, "google": 0.38, "tiktok": 0.22},
        "assist_report": [
            {"channel": "meta", "last_touch": 15, "first_touch": 20, "assists": 10,
             "assist_ratio": 0.4, "total_involvement": 25, "channel_role": "Hibrit"},
            {"channel": "google", "last_touch": 18, "first_touch": 8, "assists": 5,
             "assist_ratio": 0.22, "total_involvement": 23, "channel_role": "Dönüştürücü"},
        ],
        "insights": [
            {"type": "success", "category": "attribution", "text": "Meta en güçlü kanal."},
            {"type": "warning", "category": "data", "text": "Tek temaslı yolculuk oranı yüksek."},
        ],
    }


class TestBuildDDAPptx:
    def test_returns_presentation(self):
        result = _sample_dda_result()
        prs = build_dda_pptx(result, "Test Campaign", "2026-06-17")
        assert hasattr(prs, "slides")

    def test_slide_count(self):
        result = _sample_dda_result()
        prs = build_dda_pptx(result, "Test Campaign", "2026-06-17")
        assert len(prs.slides) == 5

    def test_no_insights_fewer_slides(self):
        result = _sample_dda_result()
        result["insights"] = []
        prs = build_dda_pptx(result, "Test Campaign", "2026-06-17")
        assert len(prs.slides) == 4

    def test_to_bytes(self):
        result = _sample_dda_result()
        prs = build_dda_pptx(result, "Test Campaign", "2026-06-17")
        buf = pptx_to_bytes(prs)
        assert isinstance(buf, BytesIO)
        assert buf.getvalue()[:4] == b"PK\x03\x04"

    def test_revenue_mode(self):
        result = _sample_dda_result()
        result["bq_summary"] = {"total_revenue": 500000}
        prs = build_dda_pptx(result, "Revenue Co", "2026-06-17", objective="revenue")
        assert len(prs.slides) == 5

    def test_empty_attribution(self):
        result = _sample_dda_result()
        result["hybrid_attribution"] = {}
        prs = build_dda_pptx(result, "Empty", "2026-06-17")
        assert len(prs.slides) >= 4
