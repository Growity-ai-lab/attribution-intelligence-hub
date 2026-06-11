"""End-to-end integration tests for Time's Hub | Attribution Intelligence.

These tests exercise the full API pipeline from HTTP request through to
model computation and response serialization.
"""

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from backend.main import app

client = TestClient(app)

SAMPLE_CSV_PATH = Path(__file__).resolve().parent.parent / "data" / "sample" / "journeys_sample.csv"


# --------------- Auth Endpoints ---------------


class TestAuthEndpoints:
    def test_login_success(self):
        r = client.post(
            "/api/auth/login",
            data={"username": "admin", "password": "test-password-for-ci"},
        )
        assert r.status_code == 200
        data = r.json()
        assert "access_token" in data
        assert data["token_type"] == "bearer"

    def test_login_wrong_password(self):
        r = client.post(
            "/api/auth/login",
            data={"username": "admin", "password": "wrong"},
        )
        assert r.status_code == 401

    def test_login_wrong_username(self):
        r = client.post(
            "/api/auth/login",
            data={"username": "nobody", "password": "test-password-for-ci"},
        )
        assert r.status_code == 401

    def test_me_with_valid_token(self, auth_headers):
        r = client.get("/api/auth/me", headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        assert data["username"] == "admin"
        assert data["role"] == "admin"

    def test_me_without_token(self):
        r = client.get("/api/auth/me")
        assert r.status_code == 401

    def test_protected_endpoint_without_token(self):
        r = client.post("/api/dda/run", json={"journeys": [], "prior_alpha": 0.5})
        assert r.status_code == 401


# --------------- Config Endpoint ---------------


class TestConfigEndpoint:
    def test_channels_list(self):
        r = client.get("/api/config/channels")
        assert r.status_code == 200
        data = r.json()
        assert "channels" in data
        assert "meta" in data["channels"]
        assert len(data["channels"]) == 10

    def test_config_contains_model_params(self):
        r = client.get("/api/config/channels")
        data = r.json()
        assert "adstock_params" in data
        assert "saturation_params" in data
        assert "unified_weights" in data
        assert "dda_blend_weights" in data
        assert data["unified_weights"]["dda"] == 0.50


# --------------- Weekly Upload Flow ---------------


class TestWeeklyUploadFlow:
    def test_upload_valid_csv(self, sample_weekly_csv, auth_headers):
        r = client.post(
            "/api/data/upload",
            files={"file": ("weekly.csv", io.BytesIO(sample_weekly_csv), "text/csv")},
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert data["rows"] == 2
        assert "meta" in data["channels"]
        assert "google" in data["channels"]

    def test_upload_returns_weeks(self, sample_weekly_csv, auth_headers):
        r = client.post(
            "/api/data/upload",
            files={"file": ("weekly.csv", io.BytesIO(sample_weekly_csv), "text/csv")},
            headers=auth_headers,
        )
        data = r.json()
        assert "2026-W06" in data["weeks"]


# --------------- MMM Endpoints ---------------


class TestMMMEndpoints:
    def test_adstock_computation(self):
        r = client.get("/api/mmm/adstock/meta?spend=1000000,500000,250000")
        assert r.status_code == 200
        data = r.json()
        assert data["channel"] == "meta"
        assert len(data["adstocked"]) == 3
        # Adstock carry-over means second value > raw spend
        assert data["adstocked"][1] > 500000

    def test_saturation_computation(self):
        r = client.get("/api/mmm/saturation/google?values=0,500000,1000000")
        assert r.status_code == 200
        data = r.json()
        assert data["channel"] == "google"
        assert len(data["saturated_values"]) == 3
        assert data["saturated_values"][0] == 0.0  # saturation(0) = 0

    def test_decomposition_returns_all_channels(self):
        r = client.get("/api/mmm/decomposition")
        assert r.status_code == 200
        data = r.json()
        assert len(data) == 10
        channels = {d["channel"] for d in data}
        assert "meta" in channels
        assert "tv_match" in channels


# --------------- DDA from CSV Flow ---------------


class TestDDAFromCSVFlow:
    def test_run_from_csv_with_fixture(self, sample_journeys_csv, auth_headers):
        r = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert "journey_stats" in data
        assert "markov" in data
        assert "hybrid_attribution" in data
        assert "unified_report" in data

    def test_conversion_count_correct(self, sample_journeys_csv, auth_headers):
        """Verify the conversion bug fix — form channel filtering must not lose conversions."""
        r = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        data = r.json()
        stats = data["journey_stats"]
        # L1 and L2 convert, L3 and L4 do not
        assert stats["converted"] == 2
        assert stats["not_converted"] == 2
        assert stats["total_journeys"] == 4

    def test_cross_validation_is_list(self, sample_journeys_csv, auth_headers):
        """Verify cross_validation serialization fix — must be a list of dicts."""
        r = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        data = r.json()
        cv = data["cross_validation"]
        assert isinstance(cv, list)
        if cv:
            assert "channel" in cv[0]
            assert "flagged" in cv[0]

    def test_markov_conversion_probability_positive(self, sample_journeys_csv, auth_headers):
        r = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        data = r.json()
        assert data["markov"]["conversion_probability"] > 0

    def test_unified_report_has_scores(self, sample_journeys_csv, auth_headers):
        r = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        data = r.json()
        report = data["unified_report"]
        assert len(report) > 0
        # DDA-only attribution: unified_score == dda_score, no MMM/incrementality
        for ch_scores in report.values():
            assert "unified_score" in ch_scores
            assert "dda_score" in ch_scores
            assert ch_scores["unified_score"] == ch_scores["dda_score"]


# --------------- DDA from JSON ---------------


class TestDDAFromJSON:
    def test_basic_run(self, minimal_journeys, auth_headers):
        r = client.post(
            "/api/dda/run",
            json={"journeys": minimal_journeys, "prior_alpha": 0.5},
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert "markov" in data
        assert "shapley_dda" in data
        assert "hybrid_attribution" in data

    def test_attribution_sums_to_one(self, minimal_journeys, auth_headers):
        r = client.post(
            "/api/dda/run",
            json={"journeys": minimal_journeys, "prior_alpha": 0.5},
            headers=auth_headers,
        )
        data = r.json()
        hybrid = data["hybrid_attribution"]
        assert sum(hybrid.values()) == pytest.approx(1.0, abs=0.01)


# --------------- Journey Stats ---------------


class TestJourneyStats:
    def test_journey_stats_endpoint(self, minimal_journeys, auth_headers):
        r = client.post("/api/dda/journey-stats", json=minimal_journeys, headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        assert data["total_journeys"] == 5
        assert data["converted"] == 3
        assert data["not_converted"] == 2
        assert 0 < data["conversion_rate"] < 1


# --------------- Reallocation Flow ---------------


class TestReallocationFlow:
    def test_dda_to_reallocation_e2e(self, sample_journeys_csv, auth_headers):
        """Full flow: CSV → DDA → unified_report → reallocation."""
        # Step 1: Run DDA
        r1 = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        assert r1.status_code == 200
        unified_report = r1.json()["unified_report"]

        # Step 2: Run reallocation
        current_budgets = {ch: 100_000 for ch in unified_report}
        r2 = client.post(
            "/api/unified/reallocation",
            json={"unified_report": unified_report, "current_budgets": current_budgets},
            headers=auth_headers,
        )
        assert r2.status_code == 200
        data = r2.json()
        total_suggested = sum(s["suggested"] for s in data["suggestions"].values())
        assert total_suggested == pytest.approx(data["total_budget"])

    def test_reallocation_higher_score_gets_more(self, sample_journeys_csv, auth_headers):
        """Channel with higher unified_score should get more budget."""
        r1 = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        unified_report = r1.json()["unified_report"]

        current_budgets = {ch: 100_000 for ch in unified_report}
        r2 = client.post(
            "/api/unified/reallocation",
            json={"unified_report": unified_report, "current_budgets": current_budgets},
            headers=auth_headers,
        )
        suggestions = r2.json()["suggestions"]

        # Find highest and lowest unified_score channels
        scores = {ch: unified_report[ch]["unified_score"] for ch in unified_report}
        top_ch = max(scores, key=scores.get)
        bottom_ch = min(scores, key=scores.get)

        if scores[top_ch] != scores[bottom_ch]:
            assert suggestions[top_ch]["suggested"] > suggestions[bottom_ch]["suggested"]


# --------------- Channel Benchmarks (Media-Planning Validation) ---------------


class TestChannelBenchmarks:
    def test_no_dda_run_returns_unavailable(self, make_campaign, auth_headers):
        """Benchmark endpoint returns available=False when no DDA run is stored."""
        campaign_id = make_campaign("Benchmark Campaign")
        r = client.get(
            f"/api/benchmarks/channel-metrics?campaign_id={campaign_id}",
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert data["available"] is False

    def test_dda_run_populates_benchmarks(self, sample_journeys_csv, make_campaign, auth_headers):
        """After a DDA run with campaign_id, benchmarks become available."""
        campaign_id = make_campaign("Benchmark Campaign")

        # Run DDA bound to the campaign — should persist a DDAResult
        r1 = client.post(
            f"/api/dda/run-from-csv?campaign_id={campaign_id}",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        assert r1.status_code == 200
        assert r1.json()["persisted"] is True

        # Benchmarks should now be available with per-channel metrics
        r2 = client.get(
            f"/api/benchmarks/channel-metrics?campaign_id={campaign_id}",
            headers=auth_headers,
        )
        assert r2.status_code == 200
        data = r2.json()
        assert data["available"] is True
        assert data["data_source"] == "csv"
        assert "channels" in data and len(data["channels"]) > 0
        # Every channel carries the benchmark signal fields
        for metrics in data["channels"].values():
            assert "dda_weight" in metrics
            assert "assist_ratio" in metrics
            assert "touchpoints" in metrics
        assert 0 <= data["overall_conversion_rate"] <= 1

    def test_benchmark_requires_auth(self):
        r = client.get("/api/benchmarks/channel-metrics?campaign_id=1")
        assert r.status_code == 401

    def test_plan_reconciliation(self, sample_journeys_csv, make_campaign, auth_headers):
        """A saved plan reconciled against a stored DDA run returns deviations."""
        campaign_id = make_campaign("Benchmark Campaign")

        # Stored DDA run for the campaign
        r1 = client.post(
            f"/api/dda/run-from-csv?campaign_id={campaign_id}",
            files={"file": ("journeys.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        assert r1.status_code == 200

        # Save a digital media plan for the 'meta' channel
        rsave = client.post(
            "/api/media-planning/save",
            json={
                "name": "Meta Plan",
                "channel": "meta",
                "weekly_grps": [1_000_000, 1_000_000],
                "response_snapshot": {
                    "summary": {
                        "total_spend": 2_000_000,
                        "total_leads": 500,
                        "total_funnel_leads": 540,
                        "avg_cpl": 4000,
                    }
                },
                "campaign_id": campaign_id,
                "mode": "digital",
            },
            headers=auth_headers,
        )
        assert rsave.status_code == 200
        plan_id = rsave.json()["id"]

        # Reconcile
        r2 = client.post(
            "/api/benchmarks/plan-reconciliation",
            json={"plan_id": plan_id},
            headers=auth_headers,
        )
        assert r2.status_code == 200
        data = r2.json()
        assert data["available"] is True
        assert data["channel"] == "meta"
        assert "planned" in data and "actual" in data and "deviations" in data
        assert data["planned"]["total_spend"] == 2_000_000
        assert "verdict" in data["deviations"]

    def test_reconciliation_unavailable_without_dda(self, make_campaign, auth_headers):
        """Reconciliation returns available=False when no DDA run is stored."""
        campaign_id = make_campaign("Benchmark Campaign")
        rsave = client.post(
            "/api/media-planning/save",
            json={
                "name": "Orphan Plan",
                "channel": "google",
                "weekly_grps": [500_000],
                "response_snapshot": {"summary": {"total_spend": 500_000, "total_leads": 100}},
                "campaign_id": campaign_id,
                "mode": "digital",
            },
            headers=auth_headers,
        )
        plan_id = rsave.json()["id"]
        r = client.post(
            "/api/benchmarks/plan-reconciliation",
            json={"plan_id": plan_id},
            headers=auth_headers,
        )
        assert r.status_code == 200
        assert r.json()["available"] is False


# --------------- Export & Insight Trends ---------------


class TestDDAExport:
    def test_export_returns_xlsx(self, sample_journeys_csv, make_campaign, auth_headers):
        """Export endpoint returns a valid xlsx file after a DDA run."""
        campaign_id = make_campaign("ExportCampaign")
        client.post(
            f"/api/dda/run-from-csv?campaign_id={campaign_id}",
            files={"file": ("j.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        r = client.get(
            f"/api/export/dda-report?campaign_id={campaign_id}",
            headers=auth_headers,
        )
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        assert len(r.content) > 500

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        assert len(wb.sheetnames) >= 3
        assert "Özet" in wb.sheetnames
        assert "Kanal Atfetme" in wb.sheetnames
        assert "Asist Raporu" in wb.sheetnames

    def test_export_404_no_dda_run(self, make_campaign, auth_headers):
        """Export returns 404 when no DDA run exists."""
        campaign_id = make_campaign("ExportCampaign")
        r = client.get(
            f"/api/export/dda-report?campaign_id={campaign_id}",
            headers=auth_headers,
        )
        assert r.status_code == 404

    def test_export_requires_auth(self):
        r = client.get("/api/export/dda-report?campaign_id=1")
        assert r.status_code == 401


class TestInsightTrends:
    def test_no_runs_returns_unavailable(self, make_campaign, auth_headers):
        campaign_id = make_campaign("TrendCampaign")
        r = client.get(
            f"/api/insights/trend?campaign_id={campaign_id}",
            headers=auth_headers,
        )
        assert r.status_code == 200
        assert r.json()["available"] is False
        assert r.json()["run_count"] == 0

    def test_single_run_returns_unavailable(self, sample_journeys_csv, make_campaign, auth_headers):
        campaign_id = make_campaign("TrendCampaign")
        client.post(
            f"/api/dda/run-from-csv?campaign_id={campaign_id}",
            files={"file": ("j.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        r = client.get(
            f"/api/insights/trend?campaign_id={campaign_id}",
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert data["available"] is False
        assert data["run_count"] == 1

    def test_two_runs_returns_comparison(self, sample_journeys_csv, make_campaign, auth_headers):
        campaign_id = make_campaign("TrendCampaign")
        for _ in range(2):
            client.post(
                f"/api/dda/run-from-csv?campaign_id={campaign_id}",
                files={"file": ("j.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
                headers=auth_headers,
            )
        r = client.get(
            f"/api/insights/trend?campaign_id={campaign_id}",
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert data["available"] is True
        assert isinstance(data["insights"], list)
        assert "current_run_date" in data
        assert "previous_run_date" in data
        assert "summary" in data

    def test_trend_requires_auth(self):
        r = client.get("/api/insights/trend?campaign_id=1")
        assert r.status_code == 401


class TestCompareSnapshots:
    def test_weight_shift_detected(self):
        from backend.models.dda.insights import compare_snapshots
        current = {
            "hybrid_attribution": {"meta": 0.30, "google": 0.50, "tiktok": 0.20},
            "journey_stats": {"total_journeys": 100, "converted": 10, "conversion_rate": 0.10, "avg_path_length": 2.0},
            "assist_report": [
                {"channel": "meta", "last_touch": 5, "first_touch": 3, "assists": 2, "assist_ratio": 0.4, "total_involvement": 10},
                {"channel": "google", "last_touch": 3, "first_touch": 4, "assists": 3, "assist_ratio": 0.5, "total_involvement": 10},
            ],
        }
        previous = {
            "hybrid_attribution": {"meta": 0.40, "google": 0.40, "tiktok": 0.20},
            "journey_stats": {"total_journeys": 90, "converted": 9, "conversion_rate": 0.10, "avg_path_length": 2.0},
            "assist_report": [
                {"channel": "meta", "last_touch": 5, "first_touch": 3, "assists": 2, "assist_ratio": 0.4, "total_involvement": 10},
                {"channel": "google", "last_touch": 3, "first_touch": 4, "assists": 3, "assist_ratio": 0.5, "total_involvement": 10},
            ],
        }
        insights = compare_snapshots(current, previous)
        categories = [i["category"] for i in insights]
        assert "attribution_shift" in categories
        shift = next(i for i in insights if i["category"] == "attribution_shift")
        assert "meta" in shift["text"].lower() or "google" in shift["text"].lower()

    def test_new_channel_detected(self):
        from backend.models.dda.insights import compare_snapshots
        current = {
            "hybrid_attribution": {"meta": 0.50, "google": 0.40, "linkedin": 0.10},
            "journey_stats": {"total_journeys": 100, "converted": 10, "conversion_rate": 0.10, "avg_path_length": 2.0},
            "assist_report": [],
        }
        previous = {
            "hybrid_attribution": {"meta": 0.55, "google": 0.45},
            "journey_stats": {"total_journeys": 100, "converted": 10, "conversion_rate": 0.10, "avg_path_length": 2.0},
            "assist_report": [],
        }
        insights = compare_snapshots(current, previous)
        categories = [i["category"] for i in insights]
        assert "channel_emergence" in categories
        emergence = next(i for i in insights if i["category"] == "channel_emergence")
        assert "linkedin" in emergence["text"].lower()

    def test_identical_snapshots_no_insights(self):
        from backend.models.dda.insights import compare_snapshots
        snap = {
            "hybrid_attribution": {"meta": 0.50, "google": 0.50},
            "journey_stats": {"total_journeys": 100, "converted": 10, "conversion_rate": 0.10, "avg_path_length": 2.0},
            "assist_report": [
                {"channel": "meta", "last_touch": 5, "first_touch": 5, "assists": 0, "assist_ratio": 0.0, "total_involvement": 10},
                {"channel": "google", "last_touch": 5, "first_touch": 5, "assists": 0, "assist_ratio": 0.0, "total_involvement": 10},
            ],
        }
        insights = compare_snapshots(snap, snap)
        assert insights == []


# --------------- Sample Data Flow ---------------


class TestSampleDataFlow:
    @pytest.mark.skipif(not SAMPLE_CSV_PATH.exists(), reason="Sample CSV not found")
    def test_sample_csv_download(self):
        r = client.get("/api/data/sample/journeys")
        assert r.status_code == 200
        assert len(r.content) > 0
        # Should be valid CSV with header
        text = r.content.decode()
        assert "lead_id" in text
        assert "channel" in text

    @pytest.mark.skipif(not SAMPLE_CSV_PATH.exists(), reason="Sample CSV not found")
    def test_full_sample_pipeline(self, auth_headers):
        """Download sample → run DDA → verify 13 conversions + unified report."""
        # Fetch sample CSV
        r1 = client.get("/api/data/sample/journeys")
        assert r1.status_code == 200
        csv_bytes = r1.content

        # Run DDA from that CSV
        r2 = client.post(
            "/api/dda/run-from-csv",
            files={"file": ("journeys_sample.csv", io.BytesIO(csv_bytes), "text/csv")},
            headers=auth_headers,
        )
        assert r2.status_code == 200
        data = r2.json()

        # Verify conversion count (13/20 in sample data)
        assert data["journey_stats"]["total_journeys"] == 20
        assert data["journey_stats"]["converted"] == 13

        # Verify Markov conversion probability is meaningful
        assert data["markov"]["conversion_probability"] > 0.3

        # Verify unified report
        assert "unified_report" in data
        report = data["unified_report"]
        assert len(report) >= 6  # At least online + offline channels

        # Verify cross_validation is a list
        assert isinstance(data["cross_validation"], list)

        # Verify hybrid attribution sums to ~1
        hybrid = data["hybrid_attribution"]
        assert sum(hybrid.values()) == pytest.approx(1.0, abs=0.01)


# --------------- Demo Auth ---------------


class TestDemoAuth:
    def test_demo_login(self):
        r = client.post("/api/auth/demo")
        assert r.status_code == 200
        data = r.json()
        assert "access_token" in data
        assert data["token_type"] == "bearer"

    def test_demo_user_can_access_me(self):
        r = client.post("/api/auth/demo")
        token = r.json()["access_token"]
        r2 = client.get("/api/auth/me", headers={"Authorization": f"Bearer {token}"})
        assert r2.status_code == 200
        assert r2.json()["role"] == "demo"


# --------------- Campaign CRUD ---------------


class TestCampaignCRUD:
    def _make_client(self, auth_headers):
        r = client.post("/api/clients", json={"name": "CRUD Co", "year": 2026}, headers=auth_headers)
        assert r.status_code == 200
        return r.json()["id"]

    def test_create_and_list_campaigns(self, auth_headers):
        cid = self._make_client(auth_headers)
        r = client.post(
            f"/api/clients/{cid}/campaigns",
            json={"name": "Camp Alpha", "budget": 1_000_000, "channels": "meta,google"},
            headers=auth_headers,
        )
        assert r.status_code == 200
        camp = r.json()
        assert camp["name"] == "Camp Alpha"

        r2 = client.get(f"/api/clients/{cid}/campaigns", headers=auth_headers)
        assert r2.status_code == 200
        names = [c["name"] for c in r2.json()]
        assert "Camp Alpha" in names

    def test_delete_campaign(self, auth_headers):
        cid = self._make_client(auth_headers)
        r = client.post(
            f"/api/clients/{cid}/campaigns",
            json={"name": "ToDelete"},
            headers=auth_headers,
        )
        camp_id = r.json()["id"]
        r2 = client.delete(f"/api/campaigns/{camp_id}", headers=auth_headers)
        assert r2.status_code == 200
        r3 = client.get(f"/api/clients/{cid}/campaigns", headers=auth_headers)
        ids = [c["id"] for c in r3.json()]
        assert camp_id not in ids

    def test_delete_campaign_requires_auth(self):
        r = client.delete("/api/campaigns/999")
        assert r.status_code == 401

    def test_delete_client(self, auth_headers):
        cid = self._make_client(auth_headers)
        r = client.delete(f"/api/clients/{cid}", headers=auth_headers)
        assert r.status_code == 200


class TestCampaignObjective:
    """Lead vs revenue objective on clients and campaigns."""

    def test_client_objective_default_and_explicit(self, auth_headers):
        r = client.post("/api/clients", json={"name": "Obj Co", "year": 2026}, headers=auth_headers)
        assert r.status_code == 200
        assert r.json()["objective"] == "lead"  # default

        r2 = client.post(
            "/api/clients",
            json={"name": "Rev Co", "year": 2026, "objective": "revenue"},
            headers=auth_headers,
        )
        assert r2.json()["objective"] == "revenue"

    def test_client_invalid_objective(self, auth_headers):
        r = client.post(
            "/api/clients",
            json={"name": "Bad Co", "year": 2026, "objective": "sales"},
            headers=auth_headers,
        )
        assert r.status_code == 400

    def test_campaign_inherits_client_objective(self, auth_headers):
        rc = client.post(
            "/api/clients",
            json={"name": "Inherit Co", "year": 2026, "objective": "revenue"},
            headers=auth_headers,
        )
        cid = rc.json()["id"]
        rp = client.post(f"/api/clients/{cid}/campaigns", json={"name": "Camp"}, headers=auth_headers)
        assert rp.status_code == 200
        assert rp.json()["objective"] == "revenue"  # inherited from client

    def test_campaign_objective_override_and_lead_value(self, auth_headers):
        rc = client.post(
            "/api/clients",
            json={"name": "Override Co", "year": 2026, "objective": "revenue"},
            headers=auth_headers,
        )
        cid = rc.json()["id"]
        rp = client.post(
            f"/api/clients/{cid}/campaigns",
            json={"name": "LeadCamp", "objective": "lead", "lead_value": 15000},
            headers=auth_headers,
        )
        assert rp.json()["objective"] == "lead"
        assert rp.json()["lead_value"] == 15000

    def test_campaign_invalid_objective(self, auth_headers):
        rc = client.post("/api/clients", json={"name": "X Co", "year": 2026}, headers=auth_headers)
        cid = rc.json()["id"]
        rp = client.post(
            f"/api/clients/{cid}/campaigns",
            json={"name": "Bad", "objective": "foo"},
            headers=auth_headers,
        )
        assert rp.status_code == 400

    def test_patch_campaign_objective(self, auth_headers):
        rc = client.post("/api/clients", json={"name": "Patch Co", "year": 2026}, headers=auth_headers)
        cid = rc.json()["id"]
        rp = client.post(f"/api/clients/{cid}/campaigns", json={"name": "P"}, headers=auth_headers)
        camp_id = rp.json()["id"]
        assert rp.json()["objective"] == "lead"

        ru = client.patch(
            f"/api/campaigns/{camp_id}",
            json={"objective": "revenue", "lead_value": 5000},
            headers=auth_headers,
        )
        assert ru.status_code == 200
        assert ru.json()["objective"] == "revenue"
        assert ru.json()["lead_value"] == 5000

    def test_patch_campaign_requires_auth(self):
        r = client.patch("/api/campaigns/999", json={"objective": "lead"})
        assert r.status_code == 401


class TestBudgetSimulationModes:
    """simulate_budget objective-aware output."""

    def test_lead_mode_simulation(self, auth_headers):
        body = {
            "channel_spends": {"google / cpc": 85000, "meta / cpc": 95000},
            "dda_weights": {"google / cpc": 0.6, "meta / cpc": 0.4},
            "total_revenue": 0,
            "total_conversions": 571,
            "objective": "lead",
            "lead_value": 15000,
        }
        r = client.post("/api/simulation/budget", json=body, headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        assert data["objective"] == "lead"
        assert data["primary_metric"] == "leads"
        g = data["current"]["channels"]["google / cpc"]
        assert g["attributed_leads"] == pytest.approx(342.6, abs=1)
        assert g["cpl"] is not None
        assert g["value_roas"] is not None  # lead_value > 0
        assert data["current"]["total_attributed_value"] == pytest.approx(571 * 15000)

    def test_lead_mode_without_revenue_ok(self, auth_headers):
        # Lead mode does not require total_revenue (CSV flow has none)
        body = {
            "channel_spends": {"google / cpc": 50000},
            "dda_weights": {"google / cpc": 1.0},
            "total_conversions": 100,
            "objective": "lead",
        }
        r = client.post("/api/simulation/budget", json=body, headers=auth_headers)
        assert r.status_code == 200
        assert r.json()["current"]["total_leads"] == 100

    def test_revenue_mode_unchanged(self, auth_headers):
        body = {
            "channel_spends": {"google / cpc": 85000},
            "dda_weights": {"google / cpc": 1.0},
            "total_revenue": 1100000,
            "total_conversions": 571,
            "objective": "revenue",
        }
        r = client.post("/api/simulation/budget", json=body, headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        assert data["objective"] == "revenue"
        assert data["current"]["channels"]["google / cpc"]["roas"] is not None
        assert data["current"]["aov"] is not None


# --------------- Data Templates ---------------


class TestDataTemplates:
    def test_weekly_template(self):
        r = client.get("/api/data/template/weekly")
        assert r.status_code == 200
        text = r.content.decode()
        assert "week" in text
        assert "channel" in text

    def test_crm_template(self):
        r = client.get("/api/data/template/crm")
        assert r.status_code == 200
        text = r.content.decode()
        assert "lead_id" in text

    def test_sales_stock_template(self):
        r = client.get("/api/data/template/sales-stock")
        assert r.status_code == 200
        text = r.content.decode()
        assert "week" in text
        assert "sales_units" in text


# --------------- MMM Fit ---------------


class TestMMMFit:
    def _make_campaign_with_data(self, auth_headers):
        rc = client.post("/api/clients", json={"name": "FitCo", "year": 2026}, headers=auth_headers)
        cid = rc.json()["id"]
        rp = client.post(f"/api/clients/{cid}/campaigns", json={"name": "FitCamp"}, headers=auth_headers)
        camp_id = rp.json()["id"]
        weeks = [f"2026-W{w:02d}" for w in range(1, 13)]
        csv_lines = ["week,channel,spend,impressions,clicks,leads"]
        for w in weeks:
            csv_lines.append(f"{w},meta,{200_000},{400_000},{8000},{100}")
            csv_lines.append(f"{w},google,{80_000},{160_000},{4800},{60}")
        csv_bytes = "\n".join(csv_lines).encode()
        client.post(
            f"/api/data/upload?campaign_id={camp_id}",
            files={"file": ("weekly.csv", io.BytesIO(csv_bytes), "text/csv")},
            headers=auth_headers,
        )
        return camp_id

    def test_fit_and_status(self, auth_headers):
        camp_id = self._make_campaign_with_data(auth_headers)
        r = client.post(f"/api/mmm/fit?campaign_id={camp_id}", headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        assert "channels" in data
        assert "fit_quality" in data

        r2 = client.get(f"/api/mmm/fit-status?campaign_id={camp_id}", headers=auth_headers)
        assert r2.status_code == 200
        status = r2.json()
        assert status["has_fit"] is True

    def test_fit_status_no_data(self, auth_headers):
        rc = client.post("/api/clients", json={"name": "EmptyFit", "year": 2026}, headers=auth_headers)
        cid = rc.json()["id"]
        rp = client.post(f"/api/clients/{cid}/campaigns", json={"name": "NoCamp"}, headers=auth_headers)
        camp_id = rp.json()["id"]
        r = client.get(f"/api/mmm/fit-status?campaign_id={camp_id}", headers=auth_headers)
        assert r.status_code == 200
        assert r.json()["has_fit"] is False

    def test_fit_requires_auth(self):
        r = client.post("/api/mmm/fit?campaign_id=1")
        assert r.status_code == 401


# --------------- Media Planning CRUD ---------------


class TestMediaPlanningCRUD:
    def test_save_and_list(self, make_campaign, auth_headers):
        camp_id = make_campaign("PlanCamp")
        plan = {
            "name": "Meta Q1",
            "channel": "meta",
            "weekly_grps": [500_000, 600_000, 700_000],
            "response_snapshot": {"summary": {"total_spend": 1_800_000}},
            "campaign_id": camp_id,
            "mode": "digital",
        }
        r = client.post("/api/media-planning/save", json=plan, headers=auth_headers)
        assert r.status_code == 200
        plan_id = r.json()["id"]

        r2 = client.get(f"/api/media-planning/saved?campaign_id={camp_id}", headers=auth_headers)
        assert r2.status_code == 200
        plans = r2.json()
        assert any(p["id"] == plan_id for p in plans)

    def test_get_by_id(self, make_campaign, auth_headers):
        camp_id = make_campaign("PlanCamp")
        plan = {
            "name": "Google Q2",
            "channel": "google",
            "weekly_grps": [300_000],
            "response_snapshot": {"summary": {"total_spend": 300_000}},
            "campaign_id": camp_id,
            "mode": "digital",
        }
        r = client.post("/api/media-planning/save", json=plan, headers=auth_headers)
        plan_id = r.json()["id"]

        r2 = client.get(f"/api/media-planning/saved/{plan_id}", headers=auth_headers)
        assert r2.status_code == 200
        data = r2.json()
        assert data["name"] == "Google Q2"
        assert data["channel"] == "google"

    def test_delete_plan(self, make_campaign, auth_headers):
        camp_id = make_campaign("PlanCamp")
        plan = {
            "name": "Temp Plan",
            "channel": "meta",
            "weekly_grps": [100_000],
            "response_snapshot": {},
            "campaign_id": camp_id,
            "mode": "digital",
        }
        r = client.post("/api/media-planning/save", json=plan, headers=auth_headers)
        plan_id = r.json()["id"]

        r2 = client.delete(f"/api/media-planning/saved/{plan_id}", headers=auth_headers)
        assert r2.status_code == 200

        r3 = client.get(f"/api/media-planning/saved/{plan_id}", headers=auth_headers)
        assert r3.status_code == 404

    def test_save_requires_auth(self):
        r = client.post("/api/media-planning/save", json={"name": "x"})
        assert r.status_code == 401

    def test_presets_endpoint(self, auth_headers):
        r = client.get("/api/media-planning/presets/tv_match", headers=auth_headers)
        assert r.status_code == 200
        data = r.json()
        assert "channel" in data


# --------------- Budget Simulation Endpoint ---------------


class TestBudgetSimulationEndpoint:
    def test_simulate_budget(self, auth_headers):
        r = client.post(
            "/api/simulation/budget",
            json={
                "channel_spends": {"meta": 200_000, "google": 100_000},
                "dda_weights": {"meta": 0.6, "google": 0.4},
                "total_revenue": 1_000_000,
                "total_conversions": 200,
            },
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert "current" in data
        assert "recommendations" in data

    def test_simulate_requires_auth(self):
        r = client.post("/api/simulation/budget", json={})
        assert r.status_code == 401


# --------------- Alerts (unit tests only — endpoints not yet implemented) ------
# Alert API endpoints are planned but not yet built. Alert rule evaluation
# is tested in tests/test_modules.py::TestAlertRules.


# --------------- Health Check ---------------


class TestHealthCheck:
    def test_health(self):
        r = client.get("/api/health")
        assert r.status_code == 200


# --------------- CPL Target Planner ---------------


class TestCplTargetPlanner:
    def test_cpl_plan_basic(self, sample_journeys_csv, auth_headers):
        cr = client.post(
            "/api/clients",
            json={"name": "CplPlanClient", "year": 2099, "objective": "lead"},
            headers=auth_headers,
        )
        client_id = cr.json()["id"]
        camp = client.post(
            f"/api/clients/{client_id}/campaigns",
            json={"name": "CplCamp", "budget": 50000, "channels": "meta,google", "lead_value": 500},
            headers=auth_headers,
        )
        camp_id = camp.json()["id"]
        dda = client.post(
            f"/api/dda/run-from-csv?campaign_id={camp_id}",
            files={"file": ("j.csv", io.BytesIO(sample_journeys_csv), "text/csv")},
            headers=auth_headers,
        )
        weights = dda.json()["hybrid_attribution"]

        r = client.post(
            "/api/simulation/cpl-target",
            json={
                "target_cpl": 500,
                "target_leads": 100,
                "channel_weights": weights,
                "current_spends": {ch: 5000 for ch in list(weights.keys())[:3]},
                "lead_value": 500,
            },
            headers=auth_headers,
        )
        assert r.status_code == 200
        data = r.json()
        assert data["total_budget"] == 50000
        assert data["target_cpl"] == 500
        assert data["target_leads"] == 100
        assert "feasibility" in data
        assert "channels" in data
        assert "confidence" in data

    def test_cpl_plan_missing_fields(self, auth_headers):
        r = client.post(
            "/api/simulation/cpl-target",
            json={"target_cpl": 500},
            headers=auth_headers,
        )
        assert r.status_code == 400

    def test_cpl_plan_requires_auth(self):
        r = client.post("/api/simulation/cpl-target", json={})
        assert r.status_code == 401


# --------------- Hill Inverse & Blended CPA ---------------


class TestHillInverseAndBlendedCPA:
    def test_hill_inverse_basic(self):
        from backend.models.mmm import compute_hill_inverse
        x = compute_hill_inverse(0.5, alpha=1000, gamma=1.5)
        assert x == pytest.approx(1000, rel=0.01)

    def test_hill_inverse_low_saturation(self):
        from backend.models.mmm import compute_hill_inverse
        x = compute_hill_inverse(0.1, alpha=1000, gamma=1.5)
        assert x < 1000
        assert x > 0

    def test_hill_inverse_high_saturation(self):
        from backend.models.mmm import compute_hill_inverse
        x = compute_hill_inverse(0.9, alpha=1000, gamma=1.5)
        assert x > 1000

    def test_hill_inverse_invalid_s(self):
        from backend.models.mmm import compute_hill_inverse
        with pytest.raises(ValueError):
            compute_hill_inverse(0.0, alpha=1000, gamma=1.5)
        with pytest.raises(ValueError):
            compute_hill_inverse(1.0, alpha=1000, gamma=1.5)

    def test_blended_cpa_is_spend_weighted(self):
        from backend.models.simulation import _blended_metric
        channels = [
            {"spend": 20000, "attributed_conversions": 5},
            {"spend": 10000, "attributed_conversions": 5},
        ]
        result = _blended_metric(channels)
        assert result == 3000.0

    def test_blended_cpa_zero_conversions(self):
        from backend.models.simulation import _blended_metric
        channels = [
            {"spend": 20000, "attributed_conversions": 0},
            {"spend": 10000, "attributed_conversions": 0},
        ]
        result = _blended_metric(channels)
        assert result is None
